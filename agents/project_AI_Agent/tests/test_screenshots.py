"""Output screenshots live in the zip's output_screenshots folder, not the .docx."""
import io
import zipfile
from unittest.mock import Mock

from app.graph import nodes, prompts, revision
from app.ingestion.screenshots import is_screenshot_path, screenshot_status
from app.ingestion.zip_ingest import ingest_zip
from PIL import Image

REQUIRED = [
    {"filename": "01-login-form.png", "module": "Login form", "description": "the login form with an invalid email error visible",
     "how_to_capture": "Run python src/main.py, submit the form empty.", "linked_requirement": "Validate the login form"},
    {"filename": "02-dashboard.png", "module": "Dashboard", "description": "the dashboard after signing in with three items listed",
     "how_to_capture": "Sign in with the demo user.", "linked_requirement": "Show the dashboard"},
    {"filename": "03-tests-passing.png", "module": "Test run", "description": "the terminal showing every test passing",
     "how_to_capture": "Run python -m pytest.", "linked_requirement": "Include tests"},
]
GUIDE = {"required_screenshots": REQUIRED, "docx_required_sections": ["Problem Statement", "Approach", "Conclusion"]}


def png_bytes(color="white"):
    buffer = io.BytesIO()
    Image.new("RGB", (40, 24), color).save(buffer, format="PNG")
    return buffer.getvalue()


def make_zip(path, files):
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return str(path)


def test_only_images_inside_a_screenshots_folder_count_as_screenshots():
    assert is_screenshot_path("Proj/output_screenshots/01-login.png")
    assert is_screenshot_path("Proj/Screenshots/a.JPG")
    assert not is_screenshot_path("Proj/assets/logo.png")            # an image, but not a screenshot
    assert not is_screenshot_path("Proj/output_screenshots/notes.txt")  # a screenshot folder, but not an image
    assert not is_screenshot_path("Proj/output_screenshots")


def test_zip_ingest_finds_real_screenshots_and_ignores_fake_and_unrelated_images(tmp_path):
    archive = make_zip(tmp_path / "p.zip", {
        "Proj/src/main.py": "print(1)\n",
        "Proj/output_screenshots/01-login-form.png": png_bytes("red"),
        "Proj/output_screenshots/02-dashboard.png": png_bytes("blue"),
        "Proj/output_screenshots/03-broken.png": b"this is not an image",
        "Proj/output_screenshots/04-empty.png": b"",
        "Proj/assets/logo.png": png_bytes("green"),
    })
    evidence = ingest_zip(archive)["screenshot_evidence"]
    assert evidence["valid_files"] == ["Proj/output_screenshots/01-login-form.png", "Proj/output_screenshots/02-dashboard.png"]
    assert "Proj/output_screenshots/03-broken.png" in evidence["files"]      # seen, but does not count
    assert "Proj/assets/logo.png" not in evidence["files"]
    assert evidence["ocr_text"]


def test_a_zip_with_no_screenshots_reports_none(tmp_path):
    evidence = ingest_zip(make_zip(tmp_path / "p.zip", {"Proj/src/main.py": "print(1)\n"}))["screenshot_evidence"]
    assert evidence["valid_files"] == [] and "No screenshot images" in evidence["ocr_text"]


def test_screenshot_status_counts_real_images_against_what_is_required():
    status = screenshot_status(GUIDE, {"valid_files": ["P/output_screenshots/01-login-form.png", "P/output_screenshots/whatever.png"],
                                       "files": ["P/output_screenshots/01-login-form.png", "P/output_screenshots/whatever.png", "P/output_screenshots/x.png"]})
    assert status["shortfall"] == 1 and status["complete"] is False
    assert [item["filename"] for item in status["matched"]] == ["01-login-form.png"]      # matched by file name
    assert [item["filename"] for item in status["unmatched"]] == ["02-dashboard.png", "03-tests-passing.png"]
    assert status["unreadable_files"] == ["P/output_screenshots/x.png"]


def test_enough_images_completes_it_whatever_they_are_called():
    status = screenshot_status(GUIDE, {"valid_files": ["a/s/one.png", "a/s/two.png", "a/s/three.png"], "files": []})
    assert status["complete"] is True and status["shortfall"] == 0


def test_no_required_screenshots_means_nothing_to_check():
    assert screenshot_status({}, None)["complete"] is True
    assert screenshot_status({"required_screenshots": []}, {"valid_files": []})["complete"] is True


def test_the_missing_screenshots_message_names_each_one_with_module_description_and_how_to_capture():
    text = revision.explain_missing_screenshots(screenshot_status(GUIDE, {"valid_files": ["P/output_screenshots/01-login-form.png"], "files": []}))
    assert "this project needs 3, and your zip has 1" in text
    assert "Already found by file name: 01-login-form.png" in text
    assert "02-dashboard.png [Dashboard]: the dashboard after signing in with three items listed." in text
    assert "How to capture it: Sign in with the demo user." in text
    assert "03-tests-passing.png [Test run]" in text
    assert "01-login-form.png [" not in text.split("Already found")[1]        # what is already there is not asked for again
    assert "NOT in the .docx report" not in text and "not in the .docx report" in text
    assert revision.explain_missing_screenshots(screenshot_status(GUIDE, {"valid_files": ["a", "b", "c"]})) is None


def test_a_zip_with_no_screenshots_folder_gets_one_clear_message_not_two():
    state = {
        "submission_guide": {**GUIDE, "required_paths": [{"path": "P/src", "type": "dir"}, {"path": "P/output_screenshots", "type": "dir"}]},
        "screenshot_evidence": {"valid_files": [], "files": []},
        "structure_score": {"is_complete": True},
        "zip_structure_score": {"required_paths_detail": {"missing_items": [{"path": "P/output_screenshots", "type": "dir", "description": "Screenshots."}]}},
    }
    items = revision.revision_items(state)
    assert len(items) == 1
    assert "this project needs 3, and your zip has 0" in items[0]
    assert "no readable screenshot images in an output_screenshots folder" in items[0]
    assert 'no "output_screenshots" folder' not in items[0]


def test_the_gate_stops_a_submission_that_is_short_of_screenshots(monkeypatch):
    base = {"structure_score": {"is_complete": True}, "submission_guide": GUIDE}
    assert nodes.structure_gate_passed({**base, "screenshot_evidence": {"valid_files": ["a.png"]}}) is False
    assert nodes.structure_gate_passed({**base, "screenshot_evidence": {"valid_files": ["a.png", "b.png", "c.png"]}}) is True
    assert nodes.structure_gate_passed({"structure_score": {"is_complete": True}, "submission_guide": {}}) is True


def test_request_revision_lists_the_missing_screenshots_specifically(monkeypatch, database):
    monkeypatch.setattr(nodes, "build_review_markdown", lambda state: "review")
    state = {"submission_id": "missing", "submission_guide": GUIDE, "screenshot_evidence": {"valid_files": [], "files": []},
             "structure_score": {"is_complete": True}, "zip_structure_score": {"is_complete": True}}
    notes = nodes.request_revision_node(state)["revision_notes"]
    assert notes.startswith("- **Output screenshots are missing")
    assert "01-login-form.png [Login form]" in notes and "How to capture it: Run python src/main.py" in notes


def test_screenshot_verdicts_from_the_model_are_normalised_never_silently_present(monkeypatch):
    monkeypatch.setattr(nodes, "call_json", lambda **kwargs: {
        "requirements_check": [],
        "screenshots_check": [
            {"screenshot": "01-login-form.png", "status": "PRESENT", "evidence": "OCR shows the form"},
            {"screenshot": "02-dashboard.png", "status": "looks good", "evidence": ""},   # unknown -> unclear
            {"status": "present"},                                                         # nameless: dropped
        ]})
    result = nodes.output_verification_node({"requirements": {}, "zip_code_files": {}, "submission_guide": GUIDE})["output_verification"]
    assert [item["status"] for item in result["screenshots_check"]] == ["present", "unclear"]


def test_the_verifier_prompt_lists_the_required_screenshots_and_what_the_zip_holds():
    prompt = prompts.output_verification_prompt({
        "requirements": {}, "zip_code_files": {"P/src/main.py": "print(1)\n"}, "submission_guide": GUIDE,
        "screenshot_evidence": {"valid_files": ["P/output_screenshots/01-login-form.png"], "files": [],
                                "ocr_text": "[Screenshot 1: P/output_screenshots/01-login-form.png]\nSign in Email invalid"},
    })
    assert "01-login-form.png [Login form]: the login form with an invalid email error visible" in prompt
    assert "Screenshot images found in the zip (1)" in prompt
    assert "Sign in Email invalid" in prompt
    assert '"screenshots_check"' in prompt


def test_the_guide_prompt_asks_for_exact_names_modules_and_capture_steps():
    prompt = prompts.submission_guide_prompt({"requirements": {"functional_requirements": ["x"]}, "chosen_topic": {"title": "Login Page"}})
    for needed in ('"filename"', '"module"', '"how_to_capture"', "output_screenshots", "NOT contain code or screenshots"):
        assert needed in prompt


def test_real_graph_sends_a_zip_short_of_screenshots_back_with_the_exact_list(state, database, tmp_path, monkeypatch):
    from app.graph.graph import submission_graph
    from docx import Document
    report = tmp_path / "report.docx"
    document = Document()
    for heading in ("Problem Statement", "Approach", "Conclusion"):
        document.add_heading(heading, level=1)
        document.add_paragraph(f"A substantive explanation for the {heading} section of the report.")
    document.save(report)
    archive = make_zip(tmp_path / "p.zip", {
        "Login/src/main.py": "def login():\n    return True\n",
        "Login/output_screenshots/01-login-form.png": png_bytes("red"),
    })
    systems = []

    def provider(**kwargs):
        systems.append(kwargs["system"])
        return {"is_complete": True, "weak_sections": [], "structure_quality": "good", "clutter_flags": [], "notes": ""}

    monkeypatch.setattr(nodes, "call_json", provider)
    monkeypatch.setattr(nodes, "call_text", Mock(side_effect=AssertionError("must not write final feedback")))
    guide = {**GUIDE, "required_paths": [{"path": "Login/src", "type": "dir"}, {"path": "Login/output_screenshots", "type": "dir"}]}
    result = submission_graph.invoke({**state, "submission_id": "s1", "submission_guide": guide,
                                      "docx_path": str(report), "zip_path": archive})
    assert result["status"] == "needs_revision" and "final_score" not in result
    notes = result["revision_notes"]
    assert "this project needs 3, and your zip has 1" in notes
    assert "Already found by file name: 01-login-form.png" in notes
    assert "02-dashboard.png [Dashboard]" in notes and "03-tests-passing.png [Test run]" in notes
    assert not any("Code Quality Reviewer" in system or "Requirements Verifier" in system for system in systems)


def test_the_qa_agent_is_told_to_explain_screenshots_in_full_from_the_project_brief():
    from app.agentic import qa_agent
    prompt = qa_agent._SYSTEM_PROMPT
    assert "call get_project_brief first" in prompt
    assert "output_screenshots folder" in prompt and "NOT put in" in prompt
    assert "exact file name" in prompt and "step by step" in prompt
    assert "never say something is not" in prompt.replace("\n", " ")
