from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import Mock

import pytest
from app import config, viva
from app.api import routes
from app.db.models import AssignmentStatus, ProjectAssignment, Submission, SubmissionStatus
from app.graph import nodes
from app.llm.client import LLMError


def upload(client, report="report.docx", source="source.zip", data=b"fixture"):
    # File-content parsing is tested by the real submission graph below;
    # these API tests isolate upload validation and state persistence.
    return client.post("/api/submission/upload", data={"thread_id": "thread"}, files={"docx_file": (report, data), "zip_file": (source, data)})

def answer(client, submission_id, index, text="Explanation"):
    return client.post("/api/invoke", json={"action": "submit_viva_answer", "payload": {"submission_id": submission_id, "question_id": index, "answer": text}})

@pytest.mark.parametrize("report,source,data,status", [("report.pdf", "source.zip", b"x", 400), ("report.docx", "source.txt", b"x", 400), ("report.docx", "source.zip", b"", 400), ("report.docx", "source.zip", b"123456", 413)])
def test_invalid_uploads_leave_no_submission(client, workflow, database, monkeypatch, report, source, data, status):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 5)
    assert upload(client, report, source, data).status_code == status
    with database() as session:
        assert session.query(Submission).count() == 0
    assert not list(config.UPLOAD_DIR.iterdir())

@pytest.mark.parametrize("condition,status", [("graded", 409), ("expired", 410), ("missing", 404)])
def test_submission_guards(client, workflow, state, database, condition, status):
    if condition == "graded": state["status"] = "graded"
    with database() as session:
        assignment = session.get(ProjectAssignment, "assignment")
        if condition == "expired": assignment.deadline_at = datetime.now(timezone.utc) - timedelta(days=1)
        if condition == "missing": session.delete(assignment)
        session.commit()
    assert upload(client).status_code == status
    with database() as session:
        assert session.query(Submission).count() == 0

def test_upload_rejected_while_a_previous_submission_is_still_processing(client, workflow, database):
    with database() as session:
        session.add(Submission(id="in-flight", assignment_id="assignment", docx_path="x", zip_path="y", status=SubmissionStatus.processing))
        session.commit()
    response = upload(client)
    assert response.status_code == 409
    with database() as session:
        # Only the pre-existing in-flight row -- no second Submission was created.
        assert session.query(Submission).count() == 1

def test_a_failed_submission_run_is_marked_error_not_left_stuck_processing(client, workflow, database, monkeypatch):
    from fastapi import HTTPException
    monkeypatch.setattr(routes, "_run_submission", Mock(side_effect=HTTPException(502, "AI service unavailable")))

    first = upload(client)
    assert first.status_code == 502
    with database() as session:
        rows = session.query(Submission).all()
        assert len(rows) == 1
        assert rows[0].status == SubmissionStatus.error

    # The failed attempt must not permanently block a real retry.
    monkeypatch.setattr(routes, "_run_submission", lambda value: {"status": "graded", "passed": True, "final_score": 85, "feedback": "Good work"})
    second = upload(client)
    assert second.status_code == 200

def test_passing_submission_persists_files_and_hides_score_until_viva(client, workflow, database):
    response = upload(client)
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "pending_viva"
    assert body["final_score"] is None
    assert body["viva_question"] == {"id": 0, "question": "Explain concept 0"}
    with database() as session:
        row = session.get(Submission, body["submission_id"])
        assert row.status == SubmissionStatus.pending_viva
        assert row.score_json["final_score"] == 85
        assert Path(row.docx_path).read_bytes() == b"fixture"
        assert Path(row.zip_path).read_bytes() == b"fixture"
        assert row.viva_answers_json == []
    assert workflow.update_state.call_args.args[1] == {"status": "pending_viva"}

@pytest.mark.parametrize("result,status", [({"status": "error", "feedback": "Invalid document"}, 422), ({"status": "needs_revision", "revision_notes": "Add tests", "passed": False, "final_score": 40}, 200)])
def test_failed_evaluation_does_not_start_viva(client, workflow, monkeypatch, result, status):
    monkeypatch.setattr(routes, "_run_submission", lambda state: result)
    viva_generator = Mock(side_effect=AssertionError("Viva must not run"))
    monkeypatch.setattr(routes, "generate_viva_questions", viva_generator)
    response = upload(client)
    assert response.status_code == status
    if status == 200:
        assert response.json()["revision_notes"] == "Add tests"
        assert response.json()["passed"] is False
    else: assert response.json()["detail"] == "Invalid document"
    assert workflow.update_state.call_args.args[1]["status"] == result["status"]
    viva_generator.assert_not_called()

@pytest.mark.parametrize("correct_count,passed", [(5, False), (6, True), (10, True)])
def test_viva_threshold_progress_persistence_and_replay(client, workflow, database, monkeypatch, correct_count, passed):
    submission_id = upload(client).json()["submission_id"]
    verdict = Mock(side_effect=[{"correct": i < correct_count, "note": "Reviewed"} for i in range(10)])
    monkeypatch.setattr(routes, "verify_viva_answer", verdict)
    assert answer(client, submission_id, 1).status_code == 409
    verdict.assert_not_called()
    for index in range(10):
        response = answer(client, submission_id, index)
        assert response.status_code == 200
        result = response.json()
        if index < 9:
            assert result["status"] == "pending_viva"
            assert result["final_score"] is None
            assert result["viva_question"]["id"] == index + 1
        assert answer(client, submission_id, index).status_code == 409
    assert verdict.call_count == 10
    assert result["passed"] is passed
    assert result["viva_score"] == correct_count
    assert result["final_score"] == 85
    assert result["status"] == ("graded" if passed else "needs_revision")
    with database() as session:
        row = session.get(Submission, submission_id)
        assert len(row.viva_answers_json) == 10
        assert row.viva_passed is passed
        assert row.status.value == result["status"]

def test_unknown_submission(client):
    assert answer(client, "missing", 0).status_code == 404

def test_unknown_topic_rejected_without_advancing_graph(client, workflow, state, monkeypatch):
    state["topic_options"] = [{"id": "one", "title": "Task tracker"}]
    advance = Mock()
    monkeypatch.setattr(routes, "_invoke_graph", advance)
    assert client.post("/api/topic/choose", json={"thread_id": "thread", "topic_id": "other"}).status_code == 400
    advance.assert_not_called()

def test_topic_selection_preserves_generated_option(client, workflow, state, monkeypatch):
    topic = state["chosen_topic"]
    state["topic_options"] = [topic]
    advance = Mock(return_value={"chosen_topic": topic, "requirements": {"features": ["Save tasks"]}})
    monkeypatch.setattr(routes, "_invoke_graph", advance)
    response = client.post("/api/topic/choose", json={"thread_id": "thread", "topic_id": "one"})
    assert response.status_code == 200
    assert response.json()["requirements"] == {"features": ["Save tasks"]}
    advance.assert_called_once_with({"chosen_topic": topic}, "thread")

def test_topic_generation_includes_prior_titles(state, database, monkeypatch):
    with database() as session:
        session.get(ProjectAssignment, "assignment").topic_json = {"title": "Previous project"}
        session.commit()
    provider = Mock(return_value={"options": [{"id": "new", "title": "New project"}]})
    monkeypatch.setattr(nodes, "call_json", provider)
    assert nodes.topic_generator_node(state)["topic_options"][0]["id"] == "new"
    assert "Previous project" in provider.call_args.kwargs["system"]

def test_generation_provider_failure_is_not_a_success(state, monkeypatch):
    monkeypatch.setattr(nodes, "call_json", Mock(side_effect=LLMError("offline")))
    with pytest.raises(LLMError): nodes.topic_generator_node(state)

def test_timer_persists_topic_and_submission_window(state, database):
    before = datetime.now(timezone.utc)
    result = nodes.timer_init_node(state)
    deadline = datetime.fromisoformat(result["deadline_at"])
    assert before + timedelta(days=config.SUBMISSION_WINDOW_DAYS) <= deadline <= datetime.now(timezone.utc) + timedelta(days=config.SUBMISSION_WINDOW_DAYS)
    with database() as session:
        assignment = session.get(ProjectAssignment, "assignment")
        assert assignment.status == AssignmentStatus.in_progress
        assert assignment.topic_json == state["chosen_topic"]
        assert assignment.deadline_at is not None

@pytest.mark.parametrize("passed", [True, False])
def test_evaluation_feedback_persists_grade_and_allows_revision(state, database, monkeypatch, passed):
    with database() as session:
        session.add(Submission(id="submission", assignment_id="assignment", docx_path="report.docx", zip_path="source.zip"))
        session.commit()
    monkeypatch.setattr(nodes, "call_text", lambda **kwargs: "Evaluation feedback")
    result = nodes.feedback_generator_node({**state, "submission_id": "submission", "passed": passed, "final_score": 80 if passed else 40})
    expected = "graded" if passed else "needs_revision"
    assert result["status"] == expected
    with database() as session:
        assert session.get(Submission, "submission").status.value == expected
        assert session.get(ProjectAssignment, "assignment").status.value == expected
        assert session.get(Submission, "submission").score_json["passed"] is passed

def test_viva_question_limit_and_numbering(monkeypatch):
    monkeypatch.setattr(viva, "call_json", lambda **kwargs: {"questions": [{"id": 90, "question": "Explain", "expected_concepts": ["concept"]}] * 12})
    questions = viva.generate_viva_questions({"title": "Task tracker"}, "local", {})
    assert [q["id"] for q in questions] == list(range(10))
    assert questions[0]["expected_concepts"] == ["concept"]


def test_real_main_graph_pauses_for_topic_and_timer(state, monkeypatch, database):
    import uuid
    from app.graph.graph import compiled_graph
    topic = state["chosen_topic"]
    provider = Mock(side_effect=[{"options": [topic]}, {"features": ["Save tasks"]}, {"sections": ["Introduction"]}])
    monkeypatch.setattr(nodes, "call_json", provider)
    thread = {"configurable": {"thread_id": uuid.uuid4().hex}}
    # No certificate/enrollment gate -- every request starts straight at
    # topic generation.
    first = compiled_graph.invoke(state, thread)
    assert first["topic_options"] == [topic]
    assert compiled_graph.get_state(thread).next == ("requirement_expansion",)
    compiled_graph.update_state(thread, {"chosen_topic": topic})
    second = compiled_graph.invoke(None, thread)
    assert second["requirements"] == {"features": ["Save tasks"]}
    with database() as session:
        # Persisted so the Q&A agent can answer a requirements doubt before
        # the timer is confirmed (and about_markdown exists) -- not just
        # kept in the LangGraph checkpoint.
        assert session.get(ProjectAssignment, "assignment").requirements_json == {"features": ["Save tasks"]}
    assert compiled_graph.get_state(thread).next == ("timer_init",)
    final = compiled_graph.invoke(None, thread)
    assert final["deadline_at"]
    guide = final["submission_guide"]
    assert guide["sections"] == ["Introduction"]  # what the model wrote is kept...
    # ...but what the report needs is decided by code: three sections, no screenshots.
    assert guide["docx_required_sections"] == ["Problem Statement", "Approach", "Conclusion"]
    assert guide["required_screenshots"] == []
    assert compiled_graph.get_state(thread).next == ()
    assert provider.call_count == 3

@pytest.mark.parametrize("failure", ["docx", "zip"])
def test_real_submission_graph_stops_before_ai_on_invalid_files(state, tmp_path, monkeypatch, failure):
    from docx import Document
    from app.graph.graph import submission_graph
    report, archive = tmp_path / "report.docx", tmp_path / "source.zip"
    if failure == "docx": report.write_bytes(b"not a document")
    else: Document().save(report)
    archive.write_bytes(b"not a zip")
    provider = Mock(side_effect=AssertionError("Invalid files must not reach AI grading"))
    monkeypatch.setattr(nodes, "call_json", provider)
    result = submission_graph.invoke({**state, "docx_path": str(report), "zip_path": str(archive)})
    assert result["status"] == "error"
    assert (".docx" if failure == "docx" else ".zip") in result["feedback"]
    provider.assert_not_called()

@pytest.mark.parametrize("complete,zip_complete,expected", [(True, False, True), (False, True, False)])
def test_structure_gate_uses_report_completeness(complete, zip_complete, expected):
    assert nodes.structure_gate_passed({"structure_score": {"is_complete": complete}, "zip_structure_score": {"is_complete": zip_complete}}) is expected


# --- syntax errors: sent back before any scoring, with structured details -----

def _report_docx(path):
    from docx import Document
    document = Document()
    for heading in ("Problem Statement", "Approach", "Conclusion"):
        document.add_heading(heading, level=1)
        document.add_paragraph(f"A substantive explanation for the {heading} section of this project report.")
    document.save(path)


def _zip_with(path, files):
    import zipfile
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def test_real_submission_graph_sends_syntax_errors_back_before_any_scoring(state, database, tmp_path, monkeypatch):
    from app.graph.graph import submission_graph
    report, archive = tmp_path / "report.docx", tmp_path / "source.zip"
    _report_docx(report)
    _zip_with(archive, {
        "Tracker/src/main.py": "def total(items):\n    return sum(items\n",
        "Tracker/src/util.py": "def fine():\n    return 1\n",
    })
    systems = []

    def provider(**kwargs):
        systems.append(kwargs["system"])
        return {"is_complete": True, "weak_sections": [], "structure_quality": "good", "clutter_flags": [], "notes": ""}

    monkeypatch.setattr(nodes, "call_json", provider)
    monkeypatch.setattr(nodes, "call_text", Mock(side_effect=AssertionError("must not write final feedback")))
    guide = {"docx_required_sections": ["Problem Statement", "Approach", "Code", "Output Screenshots", "Conclusion"],
             "required_paths": [{"path": "Tracker/src", "type": "dir"}]}
    result = submission_graph.invoke({**state, "submission_id": "s1", "submission_guide": guide,
                                      "docx_path": str(report), "zip_path": str(archive)})

    assert result["status"] == "needs_revision"
    assert result["syntax_report"]["error_count"] == 1
    error = result["syntax_report"]["errors"][0]
    assert error["path"] == "Tracker/src/main.py" and error["language"] == "Python"
    assert error["line"] and error["source_line"] and error["message"]
    assert result["revision_notes"] == ""                      # nothing else was wrong
    assert "final_score" not in result and "code_quality_score" not in result
    # Only the two structure reviewers ran; no scoring model and no sandbox run.
    assert not any("Code Quality Reviewer" in system or "Requirements Verifier" in system for system in systems)
    assert "execution_result" not in result


def test_real_submission_graph_scores_a_clean_submission(state, database, tmp_path, monkeypatch):
    from app.graph.graph import submission_graph
    report, archive = tmp_path / "report.docx", tmp_path / "source.zip"
    _report_docx(report)
    _zip_with(archive, {"Tracker/src/main.py": "def total(items):\n    return sum(items)\n"})

    def provider(**kwargs):
        system = kwargs["system"]
        if "Code Quality Reviewer" in system:
            return {"structure_score": 20, "syntax_score": 25, "maintainability_score": 20, "completeness_score": 20, "total_code_score": 85}
        if "Requirements Verifier" in system:
            return {"output_correct": True, "requirements_check": [{"requirement": "Add up items", "status": "met", "evidence": "main.total"}]}
        if "Final Score Decision" in system:
            return {"final_score": 85, "reasoning": "solid"}
        return {"is_complete": True, "weak_sections": [], "structure_quality": "good", "clutter_flags": [], "notes": ""}

    monkeypatch.setattr(nodes, "call_json", provider)
    monkeypatch.setattr(nodes, "call_text", lambda **kwargs: "Well done.")
    monkeypatch.setattr(nodes, "run_python_submission", Mock(side_effect=nodes.Judge0Unavailable("offline")))
    result = submission_graph.invoke({**state, "submission_id": "s2", "docx_path": str(report), "zip_path": str(archive),
                                      "requirements": {"functional_requirements": ["Add up items"]},
                                      "submission_guide": {"docx_required_sections": ["Problem Statement", "Approach", "Conclusion"]}})
    assert result["syntax_report"]["has_errors"] is False
    assert result["output_verification"]["requirements_check"][0]["status"] == "met"
    assert result["final_score"] == 85 and result["passed"] is True


def test_upload_response_carries_the_structured_syntax_errors(client, workflow, monkeypatch):
    errors = [{"path": "P/src/main.py", "language": "Python", "line": 2, "column": 12,
               "message": "'(' was never closed", "source_line": "    return sum(items"}]
    monkeypatch.setattr(routes, "_run_submission", lambda value: {
        "status": "needs_revision", "revision_notes": "", "syntax_report": {"errors": errors, "has_errors": True}})
    response = upload(client)
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "needs_revision"
    assert body["syntax_errors"] == errors
    assert not body["revision_notes"]


def test_upload_response_has_no_syntax_errors_field_content_when_the_code_parses(client, workflow, monkeypatch):
    monkeypatch.setattr(routes, "_run_submission", lambda value: {
        "status": "needs_revision", "revision_notes": "Report: add more detail.", "syntax_report": {"errors": [], "has_errors": False}})
    body = upload(client).json()
    assert body["syntax_errors"] is None and body["revision_notes"] == "Report: add more detail."
