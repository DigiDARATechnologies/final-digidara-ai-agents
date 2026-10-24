"""Knowledge tests: documents in, questions out, taken through an assessment link."""
import io
import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.auth.security import create_access_token
from app.knowledge import extract, generate
from app.knowledge import routes as knowledge_routes
from app.models import User
from app.organizations import routes as org_routes
from app.proctoring import routes as proctoring_routes

MATERIAL = "Photosynthesis turns light energy into chemical energy. It happens in the chloroplasts of plant cells. " * 3


def auth(user_id="learner"):
    return {"authorization": "Bearer " + create_access_token(user_id)}


@pytest.fixture
def api(database, monkeypatch):
    from app.admin import routes as admin_routes
    from app.rate_limit import limiter

    monkeypatch.setattr(admin_routes, "get_session", database)
    limiter.reset()
    with database() as session:
        for user_id in ("learner", "admin", "outsider"):
            session.add(User(id=user_id, name=user_id.title(), email=f"{user_id}@example.test", email_verified=True,
                             token_balance=1000, password_hash="x"))
        session.add(User(id="platform", name="DigiDARA", email="platform@example.test", email_verified=True,
                         token_balance=1000, password_hash="x", is_admin=True))
        session.commit()
    app = FastAPI()
    for module in (org_routes, proctoring_routes, knowledge_routes):
        app.include_router(module.router)
    with TestClient(app) as client:
        yield client


@pytest.fixture
def fake_model(monkeypatch):
    """The model, faked: writes numbered questions, grades every answer 6/10."""
    calls = []

    def complete(system, user, max_tokens):
        calls.append(user)
        if system == generate.WRITE_SYSTEM:
            mcq = int(user.split("Write exactly ")[1].split(" ")[0])
            qa = int(user.split(" and ")[1].split(" ")[0])
            n = len(calls)
            return {"questions": [
                *({"type": "mcq", "question": f"Where does photosynthesis happen? ({n}.{i})",
                   "options": ["Chloroplasts", "Nucleus", "Roots", "Mitochondria"], "correct": 0,
                   "explanation": "Chloroplasts."} for i in range(mcq)),
                *({"type": "qa", "question": f"What does photosynthesis do? ({n}.{i})",
                   "answer": "It turns light energy into chemical energy.", "explanation": "First line."} for i in range(qa)),
            ]}
        if system == generate.GRADE_SYSTEM:
            items = json.loads(user.split("\n\n", 1)[1])
            return {"grades": [{"id": item["id"], "score": 6, "feedback": "Partly right."} for item in items]}
        return {"summary": "Good effort.", "strengths": ["Cells"], "improve": ["Energy"]}

    monkeypatch.setattr(generate, "_complete", complete)
    return calls


def _org(client):
    created = client.post("/organizations/platform", headers=auth("platform"), json={
        "name": "Anna University", "kind": "college", "admin_email": "admin@example.test"})
    assert created.status_code == 201, created.text
    org = client.get("/organizations/me", headers=auth("admin")).json()["membership"]["organization"]
    client.post("/organizations/join", headers=auth(), json={"join_code": org["join_code"]})
    return org["id"]


def _ready_test(client, org_id, question_type="mixed", count=4):
    base = f"/organizations/{org_id}/knowledge-tests"
    test = client.post(base, headers=auth("admin"), json={"title": "Biology unit 1", "question_type": question_type,
                                                           "num_questions": count, "seconds_per_question": 60}).json()
    up = client.post(f"{base}/{test['id']}/documents", headers=auth("admin"),
                     files=[("files", ("notes.txt", MATERIAL.encode(), "text/plain")),
                            ("files", ("old.doc", b"binary", "application/msword"))])
    assert up.status_code == 200, up.text
    results = up.json()["results"]
    assert results[0]["ok"] and not results[1]["ok"] and "old Office" in results[1]["message"]
    generated = client.post(f"{base}/{test['id']}/generate", headers=auth("admin"))
    assert generated.status_code == 202, generated.text
    return client.get(f"{base}/{test['id']}", headers=auth("admin")).json()


def test_only_admins_build_tests(api):
    org_id = _org(api)
    base = f"/organizations/{org_id}/knowledge-tests"
    body = {"title": "Quiz", "question_type": "mcq", "num_questions": 5}
    assert api.post(base, headers=auth("learner"), json=body).status_code == 403
    assert api.post(base, headers=auth("outsider"), json=body).status_code == 403
    assert api.get(base, headers=auth("learner")).status_code == 403
    assert api.post(base, headers=auth("admin"), json=body).status_code == 201
    assert api.post(base, headers=auth("admin"), json={**body, "question_type": "essay"}).status_code == 400


def test_documents_become_questions_with_answers(api, fake_model):
    org_id = _org(api)
    test = _ready_test(api, org_id, "mixed", 4)
    assert test["status"] == "ready" and test["document_count"] == 1
    kinds = [q["qtype"] for q in test["questions"]]
    assert kinds == ["mcq", "qa", "mcq", "qa"]
    mcq = test["questions"][0]
    assert len(mcq["options"]) == 4 and mcq["options"][mcq["correct"]] == "Chloroplasts"
    assert test["questions"][1]["answer"].startswith("It turns light")
    # The admin can fix a question.
    edited = api.patch(f"/organizations/{org_id}/knowledge-tests/{test['id']}/questions/{mcq['id']}", headers=auth("admin"),
                       json={"question": "Where in the cell does photosynthesis happen?", "options": ["Chloroplast", "Nucleus", "Wall", "Vacuole"], "correct": 0})
    assert edited.status_code == 200 and edited.json()["answer"] == "Chloroplast"


def test_learner_takes_it_through_a_link_and_gets_a_report(api, fake_model, database):
    org_id = _org(api)
    test = _ready_test(api, org_id, "mixed", 4)
    link = api.post(f"/organizations/{org_id}/assessments", headers=auth("admin"), json={
        "title": "Biology quiz", "kind": "knowledge_test", "config": {"test_id": test["id"]},
        "proctoring": {"tab": True, "fullscreen": False, "copy_paste": True, "webcam": False, "face": False}})
    assert link.status_code == 201, link.text
    assert link.json()["config"]["num_questions"] == 4
    token = link.json()["path"].rsplit("/", 1)[-1]
    preview = api.get(f"/assessments/{token}", headers=auth()).json()
    assert preview["kind"] == "knowledge_test" and preview["can_start"]
    session = api.post(f"/assessments/{token}/start", headers=auth(), json={"consent": True}).json()
    assert session["agent_name"] == "knowledge_test_agent"

    attempt = api.post("/knowledge/attempts", headers=auth(), json={"session_id": session["id"]}).json()
    # Learners never receive the answers.
    assert all("correct" not in q and "answer" not in q for q in attempt["questions"])
    assert api.post("/knowledge/attempts", headers=auth("outsider"), json={"session_id": session["id"]}).status_code == 404
    q1, q2, q3, q4 = attempt["questions"]
    right = test["questions"][0]["correct"]
    send = lambda q, answer, **kw: api.post(f"/knowledge/attempts/{attempt['id']}/answers", headers=auth(),
                                           json={"question_id": q["id"], "answer": answer, "time_taken": 12, **kw})
    send(q1, str(right))
    send(q2, "Light becomes chemical energy.")
    send(q3, str((test["questions"][2]["correct"] + 1) % 4))
    send(q4, "", timed_out=True)
    # An answer cannot be changed once sent.
    send(q1, str((right + 1) % 4))
    done = api.post(f"/knowledge/attempts/{attempt['id']}/finish", headers=auth()).json()
    report = done["report"]
    assert done["status"] == "completed"
    # 1 (right MCQ) + 0.6 (graded answer) + 0 + 0 (empty, not sent to the model) out of 4.
    assert report["score"] == 1.6 and report["max"] == 4 and report["percent"] == 40.0
    assert [r["verdict"] for r in report["rows"]] == ["correct", "partial", "wrong", "wrong"]
    assert report["summary"] == "Good effort." and report["learner"] == "Learner"
    pdf = api.get(f"/knowledge/attempts/{attempt['id']}/report.pdf", headers=auth())
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    # The admin sees it in the results, with the full report and PDF.
    rows = api.get(f"/organizations/{org_id}/assessments/{link.json()['id']}/attempts", headers=auth("admin")).json()
    assert rows[0]["status"] == "completed" and rows[0]["result"]["percent"] == 40.0
    admin = api.get(f"/organizations/{org_id}/proctoring/{session['id']}/knowledge-report", headers=auth("admin"))
    assert admin.status_code == 200 and admin.json()["score"] == 1.6
    assert api.get(f"/organizations/{org_id}/proctoring/{session['id']}/knowledge-report.pdf", headers=auth("admin")).content.startswith(b"%PDF")
    assert api.get(f"/organizations/{org_id}/proctoring/{session['id']}/knowledge-report", headers=auth("learner")).status_code == 403

    # Once taken, the test's questions are fixed.
    locked = api.post(f"/organizations/{org_id}/knowledge-tests/{test['id']}/generate", headers=auth("admin"))
    assert locked.status_code == 409


def test_a_link_needs_a_ready_test(api):
    org_id = _org(api)
    draft = api.post(f"/organizations/{org_id}/knowledge-tests", headers=auth("admin"),
                     json={"title": "Empty", "question_type": "mcq", "num_questions": 3}).json()
    made = api.post(f"/organizations/{org_id}/assessments", headers=auth("admin"),
                    json={"title": "Quiz", "kind": "knowledge_test", "config": {"test_id": draft["id"]}})
    assert made.status_code == 400 and "no questions" in made.json()["detail"]
    assert api.post(f"/organizations/{org_id}/knowledge-tests/{draft['id']}/generate", headers=auth("admin")).status_code == 400


def test_documents_are_read():
    from docx import Document
    from pptx import Presentation

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Mitochondria are the powerhouse of the cell and make ATP for energy.")
    document.save(buffer)
    assert extract.extract("cell.docx", buffer.getvalue()) [1].startswith("Mitochondria")

    buffer = io.BytesIO()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Newton's laws"
    slide.placeholders[1].text = "An object stays at rest unless a force acts on it."
    deck.save(buffer)
    kind, text = extract.extract("physics.pptx", buffer.getvalue())
    assert kind == "pptx" and "[Slide 1]" in text and "force acts" in text

    with pytest.raises(extract.ExtractError, match="not supported"):
        extract.extract("virus.exe", b"MZ")
    with pytest.raises(extract.ExtractError, match="no usable content"):
        extract.extract("tiny.txt", b"hi")
