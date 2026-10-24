"""Reads the text out of an uploaded knowledge document.

PDF, Word (.docx), PowerPoint (.pptx) and plain text are read directly. An
image is read by the vision model: its text, word for word, and a
description of any diagram, chart or table in it. Old binary formats (.doc,
.ppt) and scanned PDFs without a text layer are refused with a message that
says what to upload instead.
"""
from __future__ import annotations

import base64
import io
import logging
from pathlib import PurePath

logger = logging.getLogger("orchestrator.knowledge")

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_IMAGE_BYTES = 8 * 1024 * 1024
# Text kept per document; the whole test's context is capped in generate.py.
MAX_DOC_CHARS = 200_000

KINDS = {
    ".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".txt": "txt", ".md": "txt", ".csv": "txt",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image",
}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
ACCEPTED = ", ".join(sorted({ext.lstrip(".").upper() for ext in KINDS}))


class ExtractError(Exception):
    pass


def kind_of(filename: str) -> str:
    ext = PurePath(filename or "").suffix.lower()
    if ext in (".doc", ".ppt"):
        raise ExtractError(f"{filename}: old Office files are not supported. Save it as .{ext[1:]}x and upload again.")
    if ext not in KINDS:
        raise ExtractError(f"{filename}: this file type is not supported. Upload {ACCEPTED}.")
    return KINDS[ext]


def _clean(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\r", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()[:MAX_DOC_CHARS]


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractError("This PDF is password protected. Remove the password and upload again.") from exc
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {number}]\n{text}")
            if sum(len(p) for p in pages) > MAX_DOC_CHARS:
                break
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("This PDF could not be read. It may be damaged.") from exc
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("This Word file could not be read. It may be damaged.") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("This PowerPoint file could not be read. It may be damaged.") from exc
    slides = []
    for number, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        if parts:
            slides.append(f"[Slide {number}]\n" + "\n".join(parts))
    return "\n\n".join(slides)


def _txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


IMAGE_PROMPT = (
    "This image was uploaded as study material for a test. Write out all the text in it word for word. "
    "Then describe every diagram, chart, table, formula or labelled figure in enough detail that questions "
    "could be written from your description alone. Plain text only. If the image has no educational content, "
    "reply exactly: NO CONTENT"
)


def _image(data: bytes, filename: str) -> str:
    import litellm

    from app import config

    if len(data) > MAX_IMAGE_BYTES:
        raise ExtractError(f"{filename}: images must be under 8 MB.")
    mime = IMAGE_TYPES[PurePath(filename).suffix.lower()]
    url = f"data:{mime};base64," + base64.b64encode(data).decode("ascii")
    try:
        response = litellm.completion(
            model=config.LLM_MODEL, temperature=0, max_tokens=3000,
            messages=[{"role": "user", "content": [
                {"type": "text", "text": IMAGE_PROMPT},
                {"type": "image_url", "image_url": {"url": url, "detail": "high"}},
            ]}],
        )
    except Exception as exc:
        logger.warning("knowledge image read failed for %s", filename, exc_info=True)
        raise ExtractError(f"{filename}: the image could not be read right now. Try again.") from exc
    text = (response.choices[0].message.content or "").strip()
    return "" if text.upper().startswith("NO CONTENT") else text


def extract(filename: str, data: bytes) -> tuple[str, str]:
    """(kind, text) for one uploaded file, or ExtractError with a message for the admin."""
    kind = kind_of(filename)
    if not data:
        raise ExtractError(f"{filename} is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractError(f"{filename}: files must be under 20 MB.")
    reader = {"pdf": _pdf, "docx": _docx, "pptx": _pptx, "txt": _txt}.get(kind)
    text = _clean(_image(data, filename) if kind == "image" else reader(data))
    if len(text) < 30:
        if kind == "pdf":
            raise ExtractError(f"{filename}: no text was found. If it is a scanned PDF, upload photos or screenshots of the pages instead.")
        raise ExtractError(f"{filename}: no usable content was found in it.")
    return kind, text
